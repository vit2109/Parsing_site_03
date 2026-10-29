"""Проверки справки Word и встроенных карт хронологии."""

from __future__ import annotations

import unittest
from io import BytesIO
from zipfile import ZipFile

from docx import Document
from playwright.sync_api import Error as PlaywrightError

from parsing_site.history_capture import capture_history_maps
from parsing_site.history_report import build_history_report


class HistoryReportTest(unittest.TestCase):
    def test_selected_static_and_dynamic_include_both_maps(self) -> None:
        points = [
            {"ts": "2026-10-29T10:00:00+03:00", "object_date": "2026-10-29",
             "dynamic_index": 87, "dynamic_name": "Автомобиль", "lat": 55.7,
             "lng": 37.6, "status": "blue"},
            {"ts": "2026-10-29T10:05:00+03:00", "object_date": "2026-10-29",
             "dynamic_index": 87, "dynamic_name": "Автомобиль", "lat": 55.71,
             "lng": 37.61, "status": "alert"},
        ]
        rows = [{
            "ts": points[-1]["ts"], "dynamic_index": 87,
            "dynamic_name": "Автомобиль", "static_index": 1,
            "static_name": "Пост", "distance": 1.2, "status": "alert",
        }]
        report = build_history_report(
            rows, date_from="2026-10-29", date_to="2026-10-29",
            time_from="00:00", time_to="23:59", source="real",
            max_distance_km=150, static_index=1, dynamic_query="87",
            static_scene={"static_index": 1, "static_name": "Пост",
                          "lat": 55.7, "lng": 37.6,
                          "alert_radius_km": 5, "warning_radius_km": 10,
                          "blue_radius_km": 20, "points": points},
            dynamic_points=points,
            statics=[{"static_index": 1, "static_lat": 55.7, "static_lng": 37.6}],
        )
        document = Document(BytesIO(report))
        paragraphs = "\n".join(item.text for item in document.paragraphs)
        self.assertIn("Справка по хронологии объектов", paragraphs)
        self.assertIn("Траектория динамического объекта", paragraphs)
        self.assertIn("Стационарный объект и зоны", paragraphs)
        self.assertEqual(len(document.inline_shapes), 2)
        with ZipFile(BytesIO(report)) as archive:
            self.assertEqual(len([name for name in archive.namelist()
                                  if name.startswith("word/media/")]), 2)

    def test_period_without_selected_object_has_summary(self) -> None:
        report = build_history_report(
            [{"ts": "2026-10-29T10:00:00+03:00", "dynamic_index": 87,
              "dynamic_name": "Автомобиль", "static_index": 1,
            "static_name": "Пост", "distance": 2.5, "status": "blue",
            "static_lat": 55.7, "static_lng": 37.6,
            "dynamic_lat": 55.71, "dynamic_lng": 37.61}],
            date_from="2026-10-29", date_to="2026-10-29",
            time_from="00:00", time_to="23:59", source="real",
            max_distance_km=150, static_index=None, dynamic_query="",
        )
        document = Document(BytesIO(report))
        text = "\n".join(item.text for item in document.paragraphs)
        self.assertIn("1 динамический объект и 1 стационарный объект", text)
        self.assertIn("Стационарные объекты", text)
        self.assertEqual(len(document.inline_shapes), 1)

    def test_browser_map_screenshot_is_embedded_in_report(self) -> None:
        point = {
            "ts": "2026-10-29T10:00:00+03:00", "object_date": "2026-10-29",
            "dynamic_index": 87, "dynamic_name": "Автомобиль",
            "lat": 55.71, "lng": 37.61, "status": "blue",
        }
        scene = {
            "static_index": 1, "static_name": "Пост", "lat": 55.7, "lng": 37.6,
            "alert_radius_km": 5, "warning_radius_km": 10,
            "blue_radius_km": 20, "points": [point],
        }
        try:
            images = capture_history_maps(
                static_scene=scene, dynamic_points=[point],
                statics=[{"static_index": 1, "static_name": "Пост",
                          "static_lat": 55.7, "static_lng": 37.6}],
                rows=[], map_mode="offline",
            )
        except PlaywrightError as exc:
            self.skipTest(f"Chrome недоступен: {exc}")
        self.assertEqual(set(images), {"static", "dynamic"})
        self.assertTrue(all(image.startswith(b"\x89PNG\r\n\x1a\n") for image in images.values()))
        self.assertTrue(all(len(image) > 50_000 for image in images.values()))
        report = build_history_report(
            [], date_from="2026-10-29", date_to="2026-10-29",
            time_from="00:00", time_to="23:59", source="real",
            max_distance_km=150, static_index=1, dynamic_query="87",
            static_scene=scene, dynamic_points=[point], map_images=images,
        )
        with ZipFile(BytesIO(report)) as archive:
            media = [archive.read(name) for name in archive.namelist()
                     if name.startswith("word/media/")]
        self.assertEqual(set(media), set(images.values()))
