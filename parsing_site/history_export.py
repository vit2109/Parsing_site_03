"""Формирование Excel-файла с отфильтрованной хронологией."""

from __future__ import annotations

from collections.abc import Iterable
from datetime import datetime
from io import BytesIO
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.worksheet.table import Table, TableStyleInfo

STATUS_LABELS = {
    "alert": "Красная зона",
    "warning": "Жёлтая зона",
    "blue": "Синяя зона",
    "ok": "В выбранном радиусе",
}
STATUS_FILLS = {
    "alert": PatternFill("solid", fgColor="F4CCCC"),
    "warning": PatternFill("solid", fgColor="FFF2CC"),
    "blue": PatternFill("solid", fgColor="D9EAF7"),
    "ok": PatternFill("solid", fgColor="D9EAD3"),
}
HEADERS = (
    "Дата и время",
    "Индекс стационарного",
    "Стационарный объект",
    "Широта стационарного",
    "Долгота стационарного",
    "Индекс динамического",
    "Динамический объект",
    "Широта динамического",
    "Долгота динамического",
    "Расстояние, км",
    "Статус",
)


def build_history_workbook(
    rows: Iterable[dict[str, Any]],
    *,
    date_from: str,
    date_to: str,
    static_index: int | None,
    max_distance_km: float,
    dynamic_query: str,
    source: str,
    time_from: str | None = None,
    time_to: str | None = None,
) -> bytes:
    """Возвращает XLSX с метаданными фильтра и табличной выборкой."""
    row_list = list(rows)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Хронология"
    sheet.sheet_view.showGridLines = False

    sheet["A2"] = "Хронология сближений"
    sheet["A2"].font = Font(name="Arial", size=14, bold=True, color="1F1F1F")
    sheet["A3"] = "Период"
    start_label = f"{date_from} {time_from}" if time_from else date_from
    end_label = f"{date_to} {time_to}" if time_to else date_to
    sheet["B3"] = f"{start_label} — {end_label}"
    sheet["D3"] = "Стационарный объект"
    sheet["E3"] = _static_label(row_list, static_index)
    sheet["A4"] = "Радиус, км"
    sheet["B4"] = float(max_distance_km)
    sheet["D4"] = "Динамический объект"
    sheet["E4"] = dynamic_query.strip() or "Все"
    sheet["A5"] = "Источник данных"
    sheet["B5"] = "Тестовые" if source == "fake" else "Реальные"
    sheet["D5"] = "Строк в выгрузке"
    sheet["E5"] = len(row_list)

    for cell in ("A3", "D3", "A4", "D4", "A5", "D5"):
        sheet[cell].font = Font(name="Arial", size=10, bold=True, color="44546A")
    for row in sheet.iter_rows(min_row=3, max_row=5, min_col=1, max_col=5):
        for cell in row:
            cell.alignment = Alignment(vertical="center")
            if cell.font == Font():
                cell.font = Font(name="Arial", size=10, color="1F1F1F")

    header_row = 8
    for column, value in enumerate(HEADERS, start=1):
        cell = sheet.cell(row=header_row, column=column, value=value)
        cell.fill = PatternFill("solid", fgColor="1F4E78")
        cell.font = Font(name="Arial", size=10, bold=True, color="FFFFFF")
        cell.alignment = Alignment(
            horizontal="center", vertical="center", wrap_text=True
        )

    for excel_row, item in enumerate(row_list, start=header_row + 1):
        status = str(item.get("status") or "")
        values = (
            _excel_datetime(item.get("ts")),
            item.get("static_index"),
            item.get("static_name"),
            item.get("static_lat"),
            item.get("static_lng"),
            item.get("dynamic_index"),
            item.get("dynamic_name"),
            item.get("dynamic_lat"),
            item.get("dynamic_lng"),
            item.get("distance"),
            STATUS_LABELS.get(status, status),
        )
        for column, value in enumerate(values, start=1):
            cell = sheet.cell(row=excel_row, column=column, value=value)
            cell.font = Font(name="Arial", size=10, color="1F1F1F")
            cell.alignment = Alignment(vertical="center")
        sheet.cell(excel_row, 1).number_format = "dd.mm.yyyy hh:mm:ss"
        for column in (4, 5, 8, 9):
            sheet.cell(excel_row, column).number_format = "0.000000"
        sheet.cell(excel_row, 10).number_format = "0.000"
        sheet.cell(excel_row, 11).fill = STATUS_FILLS.get(status, PatternFill())

    last_row = header_row + len(row_list)
    if row_list:
        table = Table(displayName="ChronologyTable", ref=f"A{header_row}:K{last_row}")
        table.tableStyleInfo = TableStyleInfo(
            name="TableStyleMedium2",
            showFirstColumn=False,
            showLastColumn=False,
            showRowStripes=True,
            showColumnStripes=False,
        )
        sheet.add_table(table)
    else:
        sheet.auto_filter.ref = f"A{header_row}:K{header_row}"

    sheet.freeze_panes = f"A{header_row + 1}"
    sheet.row_dimensions[header_row].height = 32
    widths = (20, 20, 28, 22, 23, 19, 28, 21, 22, 18, 24)
    for index, width in enumerate(widths, start=1):
        sheet.column_dimensions[chr(64 + index)].width = width
    sheet.print_title_rows = f"{header_row}:{header_row}"
    sheet.sheet_properties.pageSetUpPr.fitToPage = True
    sheet.page_setup.fitToWidth = 1
    sheet.page_setup.fitToHeight = 0
    sheet.sheet_properties.tabColor = "1F4E78"

    output = BytesIO()
    workbook.save(output)
    return output.getvalue()


def _excel_datetime(value: Any) -> datetime | None:
    """Преобразует ISO-время в совместимый с Excel наивный ``datetime``."""
    if not value:
        return None
    if isinstance(value, datetime):
        parsed = value
    else:
        parsed = datetime.fromisoformat(str(value))
    return parsed.replace(tzinfo=None)


def _static_label(rows: list[dict[str, Any]], static_index: int | None) -> str:
    """Формирует подпись выбранной стационарной точки для шапки отчёта."""
    if static_index is None:
        return "Все стационарные объекты"
    if rows:
        return f"#{static_index} {rows[0].get('static_name') or ''}".strip()
    return f"#{static_index}"
