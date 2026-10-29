"""Справка DOCX по сохранённой хронологии с картами маршрутов."""

from __future__ import annotations

from collections import defaultdict
from datetime import datetime
from io import BytesIO
from math import cos, isfinite, radians
from typing import Any

from docx import Document
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

ZONE_NAMES = {"alert": "Красная", "warning": "Жёлтая", "blue": "Синяя", "ok": "Вне зон"}
ROUTE_COLORS = ("#155e9a", "#8b3c88", "#17816c", "#aa5b21", "#5860b0", "#b04a63")
MAP_SIZE = (1400, 860)


def _plural(value: int, one: str, few: str, many: str) -> str:
    if value % 10 == 1 and value % 100 != 11:
        return one
    if 2 <= value % 10 <= 4 and not 12 <= value % 100 <= 14:
        return few
    return many


def _dynamic_label(index: int) -> str:
    return f"№{index}" if index >= 0 else f"ID {-index}"


def _display_time(value: str) -> str:
    return datetime.fromisoformat(value).strftime("%d.%m.%Y %H:%M:%S")


def _map_font() -> ImageFont.FreeTypeFont | ImageFont.ImageFont:
    try:
        return ImageFont.truetype("arial.ttf", 22)
    except OSError:
        return ImageFont.load_default()


def _coordinates(item: dict[str, Any], lat_key: str = "lat", lng_key: str = "lng") -> tuple[float, float] | None:
    lat, lng = item.get(lat_key), item.get(lng_key)
    if lat is None or lng is None:
        return None
    lat, lng = float(lat), float(lng)
    return (lat, lng) if isfinite(lat) and isfinite(lng) else None


def _map_base() -> tuple[Image.Image, ImageDraw.ImageDraw]:
    image = Image.new("RGB", MAP_SIZE, "#f8fafc")
    draw = ImageDraw.Draw(image)
    for x in range(100, MAP_SIZE[0], 150):
        draw.line((x, 40, x, MAP_SIZE[1] - 40), fill="#e2e8f0", width=2)
    for y in range(80, MAP_SIZE[1], 150):
        draw.line((40, y, MAP_SIZE[0] - 40, y), fill="#e2e8f0", width=2)
    draw.rectangle((35, 35, MAP_SIZE[0] - 35, MAP_SIZE[1] - 35), outline="#94a3b8", width=3)
    return image, draw


def _png(image: Image.Image) -> bytes:
    output = BytesIO()
    image.save(output, format="PNG", optimize=True)
    return output.getvalue()


def _static_map(scene: dict[str, Any]) -> bytes | None:
    center = _coordinates(scene)
    if center is None:
        return None
    points = [point for point in scene.get("points", []) if _coordinates(point) is not None]
    image, draw = _map_base()
    font = _map_font()
    width, height = MAP_SIZE
    center_x, center_y = width / 2, height / 2
    lat_factor = 111.2
    lng_factor = 111.2 * max(0.05, cos(radians(center[0])))
    radii = [float(scene.get(key) or 0) for key in (
        "blue_radius_km", "warning_radius_km", "alert_radius_km",
    )]
    coordinates = [
        ((float(point["lng"]) - center[1]) * lng_factor,
         (float(point["lat"]) - center[0]) * lat_factor)
        for point in points
    ]
    extent = max(1.0, *radii, *(max(abs(x), abs(y)) for x, y in coordinates)) * 1.12
    scale = min((width - 140) / (2 * extent), (height - 140) / (2 * extent))

    def project(lat: float, lng: float) -> tuple[int, int]:
        return (round(center_x + (lng - center[1]) * lng_factor * scale),
                round(center_y - (lat - center[0]) * lat_factor * scale))

    for radius, fill, outline in (
        (radii[0], "#e1eef9", "#77a9d4"),
        (radii[1], "#fff0c8", "#d5ab36"),
        (radii[2], "#f8d9d9", "#cc7777"),
    ):
        if radius > 0:
            pixels = radius * scale
            draw.ellipse((center_x - pixels, center_y - pixels,
                          center_x + pixels, center_y + pixels),
                         fill=fill, outline=outline, width=3)
    grouped: dict[tuple[str, int], list[dict[str, Any]]] = defaultdict(list)
    for point in points:
        grouped[(str(point["object_date"]), int(point["dynamic_index"]))].append(point)
    for index, ((_, dynamic_id), route) in enumerate(grouped.items()):
        color = ROUTE_COLORS[index % len(ROUTE_COLORS)]
        path = [project(float(point["lat"]), float(point["lng"])) for point in route]
        if len(path) > 1:
            draw.line(path, fill=color, width=4, joint="curve")
        for x, y in path:
            draw.ellipse((x - 3, y - 3, x + 3, y + 3), fill=color)
        if index < 35 and path:
            x, y = path[-1]
            draw.ellipse((x - 8, y - 8, x + 8, y + 8), fill=color, outline="white", width=2)
            draw.text((x + 10, y - 8), str(dynamic_id), fill="#1e293b", font=font)
    draw.ellipse((center_x - 9, center_y - 9, center_x + 9, center_y + 9), fill="#111827", outline="white", width=3)
    draw.text((center_x + 14, center_y + 10), f"#{scene['static_index']}", fill="#111827", font=font)
    return _png(image)


def _trajectory_map(points: list[dict[str, Any]], statics: list[dict[str, Any]]) -> bytes | None:
    route = [point for point in points if _coordinates(point) is not None]
    if not route:
        return None
    image, draw = _map_base()
    font = _map_font()
    width, height = MAP_SIZE
    center_lat = (min(float(p["lat"]) for p in route) + max(float(p["lat"]) for p in route)) / 2
    center_lng = (min(float(p["lng"]) for p in route) + max(float(p["lng"]) for p in route)) / 2
    lat_factor = 111.2
    lng_factor = 111.2 * max(0.05, cos(radians(center_lat)))
    xs = [(float(p["lng"]) - center_lng) * lng_factor for p in route]
    ys = [(float(p["lat"]) - center_lat) * lat_factor for p in route]
    extent_x, extent_y = max(1.0, *(abs(x) for x in xs)), max(1.0, *(abs(y) for y in ys))
    scale = min((width - 160) / (2 * extent_x * 1.12), (height - 160) / (2 * extent_y * 1.12))

    def project(lat: float, lng: float) -> tuple[int, int]:
        return (round(width / 2 + (lng - center_lng) * lng_factor * scale),
                round(height / 2 - (lat - center_lat) * lat_factor * scale))

    by_day: dict[str, list[tuple[int, int]]] = defaultdict(list)
    for point in route:
        by_day[str(point["object_date"])].append(project(float(point["lat"]), float(point["lng"])))
    for path in by_day.values():
        if len(path) > 1:
            draw.line(path, fill="#12629c", width=5, joint="curve")
        for x, y in path:
            draw.ellipse((x - 3, y - 3, x + 3, y + 3), fill="#12629c")
    for index, point in enumerate((route[0], route[-1])):
        x, y = project(float(point["lat"]), float(point["lng"]))
        draw.ellipse((x - 11, y - 11, x + 11, y + 11),
                     fill="#16825d" if index == 0 else "#bd3d3d", outline="white", width=3)
        draw.text((x + 14, y - 8), "1" if index == 0 else "2", fill="#111827", font=font)
    for static in statics[:100]:
        coords = _coordinates(static, "static_lat", "static_lng")
        if coords is None:
            continue
        x, y = project(*coords)
        if 55 <= x <= width - 55 and 55 <= y <= height - 55:
            draw.rectangle((x - 6, y - 6, x + 6, y + 6), fill="#202a36")
            draw.text((x + 9, y - 8), f"#{static['static_index']}", fill="#202a36", font=font)
    return _png(image)


def _overview_map(rows: list[dict[str, Any]]) -> bytes | None:
    geo_rows = [
        row for row in rows
        if _coordinates(row, "static_lat", "static_lng") is not None
        and _coordinates(row, "dynamic_lat", "dynamic_lng") is not None
    ]
    if not geo_rows:
        return None
    image, draw = _map_base()
    font = _map_font()
    width, height = MAP_SIZE
    sampled = geo_rows[:2_000]
    locations = [
        coordinate for row in sampled for coordinate in (
            _coordinates(row, "static_lat", "static_lng"),
            _coordinates(row, "dynamic_lat", "dynamic_lng"),
        ) if coordinate is not None
    ]
    center_lat = (min(lat for lat, _ in locations) + max(lat for lat, _ in locations)) / 2
    center_lng = (min(lng for _, lng in locations) + max(lng for _, lng in locations)) / 2
    lng_factor = 111.2 * max(0.05, cos(radians(center_lat)))
    xs = [(lng - center_lng) * lng_factor for _, lng in locations]
    ys = [(lat - center_lat) * 111.2 for lat, _ in locations]
    scale = min((width - 160) / (2 * max(1.0, *(abs(x) for x in xs)) * 1.12),
                (height - 160) / (2 * max(1.0, *(abs(y) for y in ys)) * 1.12))

    def project(lat: float, lng: float) -> tuple[int, int]:
        return (round(width / 2 + (lng - center_lng) * lng_factor * scale),
                round(height / 2 - (lat - center_lat) * 111.2 * scale))

    colors = {"alert": "#bd3d3d", "warning": "#c89422", "blue": "#236aab"}
    for row in sampled:
        x, y = project(float(row["dynamic_lat"]), float(row["dynamic_lng"]))
        draw.ellipse((x - 4, y - 4, x + 4, y + 4),
                     fill=colors.get(str(row.get("status")), "#8291a4"))
    statics = {int(row["static_index"]): row for row in sampled}
    for static_id, row in list(statics.items())[:100]:
        x, y = project(float(row["static_lat"]), float(row["static_lng"]))
        draw.rectangle((x - 6, y - 6, x + 6, y + 6), fill="#152536")
        draw.text((x + 8, y - 8), f"#{static_id}", fill="#152536", font=font)
    return _png(image)


def _table(document: Document, headers: tuple[str, ...], rows: list[tuple[str, ...]]) -> None:
    table = document.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Light Shading Accent 1"
    table.autofit = False
    if len(headers) == 5:
        for column, width in zip(table.columns, (1.25, 1.55, 1.85, 0.65, 1.1), strict=True):
            column.width = Inches(width)
    for cell, title in zip(table.rows[0].cells, headers, strict=True):
        cell.text = title
    table.rows[0]._tr.get_or_add_trPr().append(OxmlElement("w:tblHeader"))
    for values in rows:
        row = table.add_row()
        row._tr.get_or_add_trPr().append(OxmlElement("w:cantSplit"))
        for cell, value in zip(row.cells, values, strict=True):
            cell.text = value
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                for run in paragraph.runs:
                    run.font.size = Pt(9)


def build_history_report(
    rows: list[dict[str, Any]], *, date_from: str, date_to: str,
    time_from: str, time_to: str, source: str, max_distance_km: float,
    static_index: int | None, dynamic_query: str,
    static_scene: dict[str, Any] | None = None,
    dynamic_points: list[dict[str, Any]] | None = None,
    statics: list[dict[str, Any]] | None = None,
    map_images: dict[str, bytes] | None = None,
) -> bytes:
    """Строит справку с итогами, событиями и снимками карты хронологии."""
    map_images = map_images or {}
    document = Document()
    section = document.sections[0]
    section.page_width, section.page_height = Inches(8.5), Inches(11)
    section.left_margin = section.right_margin = Inches(0.85)
    section.top_margin = section.bottom_margin = Inches(0.75)
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10)
    normal.font.color.rgb = RGBColor(28, 36, 47)
    normal.paragraph_format.space_after = Pt(7)
    for name in ("Title", "Heading 1", "Heading 2"):
        document.styles[name].font.name = "Arial"
        document.styles[name].font.color.rgb = RGBColor(0, 0, 0)
    title_properties = document.styles["Title"]._element.get_or_add_pPr()
    title_border = title_properties.find(qn("w:pBdr"))
    if title_border is not None:
        title_properties.remove(title_border)
    document.add_paragraph("Справка по хронологии объектов", style="Title")
    document.add_paragraph(
        "Справка составлена по сохранённым опросам карты. Она показывает "
        "зафиксированные сближения за указанный период и положения объектов "
        "с географическими координатами."
    )
    period = f"{date_from} {time_from} — {date_to} {time_to}"
    document.add_paragraph(f"Период: {period}")
    document.add_paragraph(
        f"Источник: {'тестовые' if source == 'fake' else 'реальные'} данные. "
        f"Радиус отбора событий: {max_distance_km:g} км. "
        f"Стационарный объект: {'#' + str(static_index) if static_index else 'все'}. "
        f"Динамический объект: {dynamic_query.strip() or 'все'}."
    )
    ids = {int(row["dynamic_index"]) for row in rows}
    stations = {int(row["static_index"]) for row in rows}
    minimum = min((float(row["distance"]) for row in rows), default=None)
    document.add_heading("Итоги", level=1)
    document.add_paragraph(
        f"В пределах заданного радиуса найдено {len(rows)} "
        f"{_plural(len(rows), 'запись', 'записи', 'записей')}, "
        f"{len(ids)} {_plural(len(ids), 'динамический объект', 'динамических объекта', 'динамических объектов')} "
        f"и {len(stations)} {_plural(len(stations), 'стационарный объект', 'стационарных объекта', 'стационарных объектов')}. "
        + (f"Минимальное расстояние — {minimum:.3f} км." if minimum is not None else "Сближения не зафиксированы.")
    )
    if len(rows) >= 20_000:
        document.add_paragraph("Выборка достигла предела 20 000 записей. Для полной справки сузьте период.")

    if static_scene is not None:
        document.add_heading("Стационарный объект и зоны", level=1)
        name = str(static_scene.get("static_name") or f"#{static_index}")
        document.add_paragraph(f"{name} (#{static_scene['static_index']}).")
        document.add_paragraph(
            "Границы зон: красная — до "
            f"{float(static_scene.get('alert_radius_km') or 0):g} км, "
            f"жёлтая — до {float(static_scene.get('warning_radius_km') or 0):g} км, "
            f"синяя — до {float(static_scene.get('blue_radius_km') or 0):g} км."
        )
        counts = {zone: set() for zone in ("alert", "warning", "blue")}
        for point in static_scene.get("points", []):
            if point.get("status") in counts:
                counts[point["status"]].add(point["dynamic_index"])
        document.add_paragraph(
            f"Среди показанных траекторий объектов по зонам: красная — {len(counts['alert'])}, "
            f"жёлтая — {len(counts['warning'])}, синяя — {len(counts['blue'])}. "
            "Один объект может встречаться в нескольких зонах."
        )
        picture = map_images.get("static") or _static_map(static_scene)
        if picture:
            document.add_picture(BytesIO(picture), width=Inches(6.65))
            document.add_paragraph(
                "Карта хронологии: круги показывают зоны стационарного объекта, "
                "линии — сохранённые траектории динамических объектов."
            )

    if dynamic_points is not None:
        document.add_heading("Траектория динамического объекта", level=1)
        if dynamic_points:
            first, last = dynamic_points[0], dynamic_points[-1]
            document.add_paragraph(
                f"Объект {_dynamic_label(int(first['dynamic_index']))} {first['dynamic_name']}. "
                f"Сохранено {len(dynamic_points)} "
                f"{_plural(len(dynamic_points), 'точка', 'точки', 'точек')}: "
                f"с {_display_time(str(first['ts']))} по {_display_time(str(last['ts']))}."
            )
            picture = map_images.get("dynamic") or _trajectory_map(dynamic_points, statics or [])
            if picture:
                document.add_picture(BytesIO(picture), width=Inches(6.65))
                document.add_paragraph(
                    "Карта хронологии: линия показывает сохранённую траекторию; "
                    "отмечены стационарные объекты."
                )
        else:
            document.add_paragraph("За выбранный период географические точки маршрута не найдены.")

    if static_index is None and dynamic_points is None and rows:
        document.add_heading("Стационарные объекты", level=1)
        summary: dict[int, dict[str, Any]] = {}
        for row in rows:
            item = summary.setdefault(int(row["static_index"]), {
                "name": row["static_name"], "events": 0, "minimum": float(row["distance"]),
            })
            item["events"] += 1
            item["minimum"] = min(item["minimum"], float(row["distance"]))
        _table(document, ("Объект", "Записей", "Минимум км"), [
            (f"#{index} {item['name']}", str(item["events"]), f"{item['minimum']:.3f}")
            for index, item in sorted(summary.items(), key=lambda pair: -pair[1]["events"])[:30]
        ])
        picture = map_images.get("overview") or _overview_map(rows)
        if picture:
            document.add_picture(BytesIO(picture), width=Inches(6.65))
            document.add_paragraph(
                "Карта хронологии: отмечены стационарные объекты, "
                "цветные точки — положения динамических объектов в соответствующих зонах."
            )

    document.add_heading("Последние события", level=1)
    if rows:
        _table(document, ("Дата и время", "Стационарный", "Динамический", "Км", "Зона"), [
            (str(row["ts"]).replace("T", " ")[:19],
             f"#{row['static_index']} {row['static_name']}",
             f"{_dynamic_label(int(row['dynamic_index']))} {row['dynamic_name']}",
             f"{float(row['distance']):.3f}", ZONE_NAMES.get(str(row["status"]), "Вне зон"))
            for row in rows[:12]
        ])
        if len(rows) > 12:
            document.add_paragraph("Показаны 12 последних записей. Полная выборка доступна через экспорт в Excel.")
    else:
        document.add_paragraph("За выбранный период событий в заданном радиусе нет.")
    output = BytesIO()
    document.save(output)
    return output.getvalue()
